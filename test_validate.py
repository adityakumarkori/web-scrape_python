import os
import sys
from decimal import Decimal
from datetime import datetime

os.environ['MONGODB_URI'] = 'mongodb://localhost:27017/'
os.environ['MONGODB_DATABASE'] = 'granit_blinkit_data'

from exporter import export_granit, export_blinkit, export_dataset
from profiles import get_profile
from unittest.mock import patch, MagicMock

# Mock MongoDB connection
with patch('database.MongoClient') as mock_client:
    mock_instance = MagicMock()
    mock_client.return_value = mock_instance
    mock_instance.admin.command.return_value = {'ok': 1}
    
    # Test GRANIT export in detail
    print("=" * 60)
    print("GRANIT Workbook - Detailed Column/Content Validation")
    print("=" * 60)
    
    granit_profile = get_profile("granit")
    
    products_granit = [
        {
            'source_id': 'granit:ART-001',
            'product_name': 'Brake Pad Set',
            'article_number': 'ART-001',
            'description': 'Front brake pads for BMW 3 Series',
            'brand': 'TestBrand',
            'category': 'Brakes',
            'ean_gtin': '4000000000006',
            'product_url': 'https://granit-parts.com/product/12345',
            'image_url': 'https://granit-parts.com/images/12345.jpg',
            'regional_pricing': [
                {'country': 'Germany', 'country_code': 'Germany', 'currency': 'EUR', 'net_price': Decimal('29.99'), 'vat_rate': Decimal('19'), 'gross_price': Decimal('35.69'), 'price_source': 'json_ld', 'effective_date': '2024-01-15'}
            ],
            'specifications': [{'specification_name': 'Width', 'specification_value': '150', 'unit': 'mm', 'compatibility': 'BMW 3 Series'}],
            'compliance': [{'regulation': 'ECE R90', 'compliance_standard': 'ECE', 'certification': 'E1', 'compliance_status': 'Compliant'}],
            'inventory': [{'warehouse': 'Central DE', 'country': 'Germany', 'available_quantity': '50', 'stock_status': 'In Stock'}],
            'source': 'granit',
            'collected_at': datetime.utcnow().isoformat() + 'Z',
            'dataset': 'granit',
            '_verified': True,
            '_synthetic': False
        }
    ]
    
    output_dir = 'output'
    os.makedirs(output_dir, exist_ok=True)
    
    result = export_dataset(products_granit, "granit", output_dir)
    
    import openpyxl
    wb = openpyxl.load_workbook(result["granit"])
    
    # ===== Products worksheet =====
    print("\n=== Products worksheet ===")
    ws = wb["Products"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Headers: {headers}")
    for r in range(2, ws.max_row + 1):
        row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"Row {r}: {row_data}")
    
    # ===== Regional Pricing worksheet =====
    print("\n=== Regional Pricing worksheet ===")
    ws = wb["Regional Pricing"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Headers: {headers}")
    for r in range(2, ws.max_row + 1):
        row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        # Check number formats
        formats = [ws.cell(row=r, column=c).number_format for c in range(1, ws.max_column + 1)]
        print(f"Row {r}: {row_data}")
        print(f"  Number formats: {formats}")
        # Check if currency/percentage formats applied
        if formats[4] != 'General':  # Net Price column
            print(f"  Net Price format: {formats[4]} (should be currency)")
        if formats[5] != 'General':  # VAT Rate column
            print(f"  VAT Rate format: {formats[5]} (should be percentage)")
        if formats[6] != 'General':  # Gross Price column
            print(f"  Gross Price format: {formats[6]} (should be currency)")
    
    # ===== Data Dictionary worksheet =====
    print("\n=== Data Dictionary worksheet ===")
    ws = wb["Data Dictionary"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Headers: {headers}")
    for r in range(2, min(r+10, ws.max_row + 1)):
        row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"Row {r}: {row_data}")
    
    # ===== Sources & Assumptions worksheet =====
    print("\n=== Sources & Assumptions worksheet ===")
    ws = wb["Sources & Assumptions"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Headers: {headers}")
    for r in range(2, ws.max_row + 1):
        row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"Row {r}: {row_data}")
    
    # ===== README worksheet =====
    print("\n=== README worksheet ===")
    ws = wb["README"]
    for r in range(1, min(15, ws.max_row + 1)):
        row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"Row {r}: {row_data}")
    
    print("\n=== GRANIT Validation Complete ===")
    
    # Now test Blinkit
    print("\n" + "=" * 60)
    print("Blinkit Workbook - Detailed Column/Content Validation")
    print("=" * 60)
    
    blinkit_profile = get_profile("blinkit")
    
    products_blinkit = [
        {
            'source_id': 'blinkit:SKU-001',
            'product_name': 'Organic Apples 1kg',
            'product_id': 'APL-001',
            'sku': 'SKU-001',
            'brand': 'FarmFresh',
            'category': 'Fruits',
            'subcategory': 'Apples',
            'product_description': 'Fresh organic apples from local farms',
            'barcode': '8901234567890',
            'product_url': 'https://blinkit.com/product/apples-1kg',
            'image_url': 'https://blinkit.com/apples.jpg',
            'pricing': [{'mrp': Decimal('150.00'), 'selling_price': Decimal('120.00'), 'discount_amount': Decimal('30.00'), 'discount_percentage': Decimal('20'), 'gst_rate': Decimal('5'), 'hsn_code': '08081000', 'price_source': 'json_ld', 'effective_date': '2024-01-15'}],
            'packaging': [{'net_quantity': Decimal('1'), 'unit': 'kg', 'package_type': 'Bag', 'ingredients': 'Organic apples', 'manufacturing_date': '2024-01-10', 'expiry_date': '2024-02-10', 'shelf_life': '30'}],
            'supplier': [{'supplier_name': 'Local Farms Co', 'supplier_location': 'Bangalore Rural', 'procurement_type': 'Direct', 'supplier_status': 'Active'}],
            'inventory': [{'dark_store_id': 'DS-001', 'dark_store_name': 'Koramangala', 'zone': 'South', 'pincode': '560034', 'available_quantity': '100', 'stock_status': 'In Stock', 'last_restocked': '2024-01-14'}],
            'delivery': [{'zone': 'South', 'pincode': '560034', 'delivery_sla': Decimal('30'), 'serviceability_status': 'Serviceable', 'effective_date': '2024-01-15'}],
            'source': 'blinkit',
            'collected_at': datetime.utcnow().isoformat() + 'Z',
            'dataset': 'blinkit',
            '_verified': True,
            '_synthetic': False
        }
    ]
    
    result = export_dataset(products_blinkit, "blinkit", output_dir)
    wb = openpyxl.load_workbook(result["blinkit"])
    
    # ===== Products worksheet =====
    print("\n=== Products worksheet ===")
    ws = wb["Products"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Headers: {headers}")
    for r in range(2, ws.max_row + 1):
        row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"Row {r}: {row_data}")
    
    # ===== Pricing & Tax worksheet =====
    print("\n=== Pricing & Tax worksheet ===")
    ws = wb["Pricing & Tax"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Headers: {headers}")
    for r in range(2, ws.max_row + 1):
        row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        formats = [ws.cell(row=r, column=c).number_format for c in range(1, ws.max_column + 1)]
        print(f"Row {r}: {row_data}")
        print(f"  Number formats: {formats}")
        # Check currency formats (columns 2=MRP, 3=Selling Price)
        if formats[1] != 'General':  # MRP
            print(f"  MRP format: {formats[1]} (should be currency)")
        if formats[2] != 'General':  # Selling Price
            print(f"  Selling Price format: {formats[2]} (should be currency)")
        # Check percentage formats (columns 5=Discount %, 6=GST Rate)
        if formats[4] != 'General':  # Discount %
            print(f"  Discount % format: {formats[4]} (should be percentage)")
        if formats[5] != 'General':  # GST Rate
            print(f"  GST Rate format: {formats[5]} (should be percentage)")
    
    # ===== Data Dictionary worksheet =====
    print("\n=== Data Dictionary worksheet ===")
    ws = wb["Data Dictionary"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Headers: {headers}")
    for r in range(2, min(8, ws.max_row + 1)):
        row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"Row {r}: {row_data}")
    
    # ===== README worksheet =====
    print("\n=== README worksheet ===")
    ws = wb["README"]
    for r in range(1, min(15, ws.max_row + 1)):
        row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"Row {r}: {row_data}")
    
    print("\n=== Blinkit Validation Complete ===")